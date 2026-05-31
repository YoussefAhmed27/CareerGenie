import io
import markdown
from xhtml2pdf import pisa
from docx import Document
from docx.shared import Pt
import re

def export_to_pdf(md_text: str) -> bytes:
    """
    Converts Markdown text to a PDF file in-memory using xhtml2pdf.
    """
    if not md_text:
        return b""
        
    # Convert Markdown to HTML
    html_content = markdown.markdown(md_text)
    
    # Add basic styling to make the PDF look like a CV
    styled_html = f"""
    <html>
    <head>
    <style>
        @page {{
            size: A4;
            margin: 2cm;
        }}
        body {{
            font-family: Helvetica, Arial, sans-serif;
            font-size: 11pt;
            line-height: 1.4;
            color: #333333;
        }}
        h1 {{ font-size: 20pt; text-align: center; color: #111111; margin-bottom: 5px; }}
        h2 {{ font-size: 14pt; color: #222222; border-bottom: 1px solid #cccccc; padding-bottom: 3px; margin-top: 15px; margin-bottom: 10px; }}
        h3 {{ font-size: 12pt; color: #333333; margin-top: 10px; margin-bottom: 5px; }}
        p {{ margin-bottom: 8px; }}
        ul {{ margin-top: 0px; margin-bottom: 10px; padding-left: 20px; }}
        li {{ margin-bottom: 4px; }}
       hr {{ border: 0; border-bottom: 1px solid #dddddd; margin: 10px 0; }}
    </style>
    </head>
    <body>
        {html_content}
    </body>
    </html>
    """
    
    pdf_buffer = io.BytesIO()
    
    # Create PDF
    pisa_status = pisa.CreatePDF(styled_html, dest=pdf_buffer)
    
    if pisa_status.err:
        raise Exception("Failed to generate PDF from Markdown")
        
    return pdf_buffer.getvalue()

def export_to_docx(md_text: str) -> bytes:
    document = Document()
    
    style = document.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    
    lines = md_text.split('\n')
    
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
            
        if stripped == '---':
            paragraph = document.add_paragraph()
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            pPr = paragraph._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'CCCCCC')
            pBdr.append(bottom)
            pPr.append(pBdr)
            continue
            
        if stripped.startswith('# '):
            title_text = stripped[2:].strip()
            title = document.add_heading(title_text, level=0)
            title.alignment = 1
        elif stripped.startswith('## '):
            document.add_heading(stripped[3:].strip(), level=1)
        elif stripped.startswith('### '):
            document.add_heading(stripped[4:].strip(), level=2)
        elif stripped.startswith('<div class="date-line">') or stripped.startswith('<table'):
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            import re as _re
            
            spans = _re.findall(r'<td[^>]*>(.*?)</td>', stripped)
            if not spans:
                spans = _re.findall(r'<span>(.*?)</span>', stripped)
            left = spans[0] if len(spans) > 0 else ''
            right = spans[1] if len(spans) > 1 else ''
            
            p = document.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            tabs = OxmlElement('w:tabs')
            tab = OxmlElement('w:tab')
            tab.set(qn('w:val'), 'right')
            tab.set(qn('w:pos'), '9360')
            tabs.append(tab)
            pPr.append(tabs)
            
            run1 = p.add_run(left)
            run1.italic = True
            p.add_run('\t')
            run2 = p.add_run(right)
            run2.italic = True
        elif stripped.startswith('- '):
            bullet_text = stripped[2:].strip()
            p = document.add_paragraph(style='List Bullet')
            _add_formatted_text(p, bullet_text)
        else:
            p = document.add_paragraph()
            _add_formatted_text(p, stripped)
            
    docx_buffer = io.BytesIO()
    document.save(docx_buffer)
    return docx_buffer.getvalue()
def _add_formatted_text(paragraph, text: str):
    """Helper to parse **bold** tags and add them to a docx paragraph."""
    # Split text by **
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            # It's bold
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)
